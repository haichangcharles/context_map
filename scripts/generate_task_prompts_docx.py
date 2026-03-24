from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
output_dir = ROOT / "output" / "doc"
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / "IRB-Task-Prompts.docx"


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"].font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run = title.add_run("Task Prompts")
run.bold = True
run.font.name = "Times New Roman"
run.font.size = Pt(14)

meta = doc.add_paragraph()
meta.add_run("Study Title: ").bold = True
meta.add_run("Mixed-Initiative Context: A User Study of Context Management for User-AI Conversations")

purpose = doc.add_paragraph()
purpose.add_run("Document Purpose: ").bold = True
purpose.add_run("This document provides the task instructions shown to participants during the study.")

doc.add_heading("General Instructions to Participants", level=1)
doc.add_paragraph(
    "Participants complete two short brainstorming tasks with AI tools. Each participant completes "
    "one task in the standard linear chat condition and the other task in the prototype condition."
)
doc.add_paragraph(
    "Participants may use the assigned system in whatever way feels natural to them. The researcher "
    "does not require participants to use specific features, follow a fixed sequence of actions, or "
    "adopt any particular strategy for organizing ideas. The goal is for participants to use the system "
    "freely while working on the task."
)
doc.add_paragraph(
    "For the prototype condition, participants receive a brief warm-up on the basic interface mechanics "
    "only. The warm-up does not tell participants how they should complete the task or how they should "
    "organize their work."
)

doc.add_heading("Task 1: Hardware Concept for Focus With Important-Message Awareness", level=1)
doc.add_paragraph(
    "Many people want to stay focused while working or studying without being interrupted by frequent "
    "notifications. At the same time, they worry about missing important information such as urgent "
    "messages, important emails, or meeting reminders. Please design a new hardware product concept "
    "that helps users stay focused while still noticing important information."
)
doc.add_paragraph(
    "You may brainstorm with the AI, ask about related concepts, technical feasibility, interaction "
    "ideas, or similar products. As the task progresses, some ideas may become less suitable. Continue "
    "developing your ideas and eventually settle on one main product concept."
)

doc.add_heading("Task 2: Hardware Concept for Managing Important Everyday Items", level=1)
doc.add_paragraph(
    "Many people carry small but important items every day, such as keys, wallets, earbuds, ID cards, "
    "or water bottles. They want to reduce the hassle of forgetting, losing, or repeatedly checking "
    "these items, but they do not want to add too much burden or rely too heavily on a phone. Please "
    "design a new hardware product concept that helps users manage and track these important everyday items."
)
doc.add_paragraph(
    "You may brainstorm with the AI, ask about related concepts, technical feasibility, interaction "
    "ideas, or similar products. As the task progresses, some ideas may become less suitable. Continue "
    "developing your ideas and eventually settle on one main product concept."
)

doc.add_heading("Warm-Up for the Prototype Condition", level=1)
for item in [
    "Before the prototype task, participants receive a brief warm-up that explains the basic interface mechanics only.",
    "The warm-up does not suggest any task strategy.",
    "The warm-up does not teach participants how they should organize their ideas.",
]:
    add_bullet(doc, item)

doc.save(output_path)
print(output_path)
