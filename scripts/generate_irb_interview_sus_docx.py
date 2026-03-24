from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
source_path = ROOT / "Study" / "19-irb-interview-and-sus-materials.md"
output_dir = ROOT / "output" / "doc"
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / "IRB-Interview-and-Oral-SUS-Materials.docx"


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def add_number(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Number")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"].font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run = title.add_run("Interview Guide")
run.bold = True
run.font.name = "Times New Roman"
run.font.size = Pt(14)

meta = doc.add_paragraph()
meta.add_run("Study Title: ").bold = True
meta.add_run("Mixed-Initiative Context: A User Study of Context Management for User-AI Conversations")

purpose = doc.add_paragraph()
purpose.add_run("Document Purpose: ").bold = True
purpose.add_run(
    "This document is provided for IRB review as the moderator-facing guide for the post-task interview. "
    "The interview includes open-ended questions and a short set of structured rating questions after "
    "participants complete the study tasks."
)

doc.add_heading("Administration Format", level=1)
for item in [
    "The moderator conducts a semi-structured interview after both task conditions are completed.",
    "The moderator then asks a short set of structured rating questions orally.",
    "For the rating portion, the moderator reads each item aloud and records the participant's verbal response on a 5-point response scale.",
    "Participants are not required to complete a separate written questionnaire unless the study procedure changes.",
]:
    add_bullet(doc, item)

doc.add_heading("Response Scale for the Structured Rating Questions", level=1)
for item in [
    "1 = Strongly disagree",
    "2 = Disagree",
    "3 = Neither agree nor disagree",
    "4 = Agree",
    "5 = Strongly agree",
]:
    add_bullet(doc, item)

doc.add_heading("Section 1. Semi-Structured Interview Guide", level=1)
doc.add_paragraph("Interview goals")
for item in [
    "RQ1: Understand what context-management strategies participants used, when they used them, and why.",
    "RQ2: Understand how participants perceived AI involvement in organizing context, including when AI initiative was helpful, unhelpful, or needed redesign.",
    "RQ3: Understand how the prototype changed workflow, effort, clarity, control, and overall collaboration compared with linear chat.",
]:
    add_bullet(doc, item)

doc.add_paragraph("Primary interview questions")
questions = [
    "Please walk me through how you organized your thinking in this system. What kinds of actions did you take to manage information or keep track of different ideas?",
    "At what moments did you feel a need to do something to the context, for example separate ideas, return to an earlier path, remove something, reorganize information, or narrow the focus?",
    "What were you trying to achieve when you made those moves? For example, were you trying to reduce confusion, compare alternatives, keep useful information active, or discard an idea that no longer fit?",
    "Were there moments when you wanted to manage the context differently but did not? What stopped you?",
    "Compared with the standard linear chat condition, did you use different strategies here for keeping track of ideas or deciding what should remain relevant? If so, how?",
    "In the prototype system, how did you interpret the AI's role? Did it feel like the AI was only generating content, or also participating in organizing the conversation?",
    "Can you recall a moment when an AI-generated structure, suggestion, or intervention influenced what you did next? What was helpful or unhelpful about it?",
    "When do you think AI initiative in organizing context is useful, and when is it not useful?",
    "If the AI introduces a new branch, suggests reorganizing information, or tries to help decide what should stay active, what kind of behavior would feel appropriate to you?",
    "If you could redesign the AI's role in this system, how proactive or restrained should it be, and what should it do differently?",
    "Compared with the linear chat condition, how did this workflow feel different in the way you moved through the task?",
    "Did this system make the task feel clearer or more confusing? Easier or more effortful? More controlled or less controlled? Please explain with examples.",
    "Did the system change how aware you were of what information was currently active or relevant? If so, how?",
    "Do you think this way of working improved your overall experience, made it worse, or changed it in a more mixed way? Why?",
    "In what kinds of real tasks would you want to use a system like this, and in what kinds of tasks would you prefer not to use it?",
]
for item in questions:
    add_number(doc, item)

doc.add_paragraph("Optional follow-up prompts")
for item in [
    "What did you do when you felt there were too many ideas or too much information on screen?",
    "Was there a moment when you intentionally kept some information active or intentionally pushed some information out of focus?",
    "Was there a moment when you returned to an earlier idea or path? Why at that point?",
    "Did the structure ever help you compare alternatives or converge on a final direction?",
    "Did the system ever make you feel you had to do extra organizational work? If yes, what kind?",
    "Did the AI ever interrupt your flow or push the conversation in a direction you did not want?",
    "Did you ever want the AI to step in more strongly? Did you ever want it to stay out of the way?",
    "Which aspects of the experience felt better because of the system, and which felt worse because of the system?",
]:
    add_bullet(doc, item)

doc.add_heading("Section 2. Structured Rating Questions Asked During the Interview", level=1)
doc.add_paragraph("Moderator instructions")
for item in [
    "Introduce this as a short set of usability questions about the prototype system.",
    "Read each statement aloud exactly as written.",
    "Ask the participant to answer using the 1-5 response scale.",
    "Record one response for each item.",
]:
    add_bullet(doc, item)

script_p = doc.add_paragraph()
script_p.add_run("Script to introduce the rating questions: ").bold = True
script_p.add_run(
    '"I am now going to read a short set of statements about the prototype system. '
    "After each statement, please respond using a scale from 1 to 5, where 1 means "
    'strongly disagree and 5 means strongly agree."'
)

doc.add_paragraph("Usability rating items")
sus_items = [
    "I think that I would like to use this system frequently.",
    "I found the system unnecessarily complex.",
    "I thought the system was easy to use.",
    "I think that I would need the support of a technical person to be able to use this system.",
    "I found the various functions in this system were well integrated.",
    "I thought there was too much inconsistency in this system.",
    "I would imagine that most people would learn to use this system very quickly.",
    "I found the system very cumbersome to use.",
    "I felt very confident using the system.",
    "I needed to learn a lot of things before I could get going with this system.",
]
for item in sus_items:
    add_number(doc, item)

doc.add_paragraph("Researcher note")
for item in [
    "These items are being administered orally as part of the interview session.",
    "For odd-numbered items, subtract 1 from the participant's score.",
    "For even-numbered items, subtract the participant's score from 5.",
    "Sum the converted item scores and multiply the total by 2.5.",
    "The resulting usability score ranges from 0 to 100.",
]:
    add_bullet(doc, item)

doc.save(output_path)
print(output_path)
